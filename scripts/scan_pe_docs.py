import docx

def scan_pe_doc(path):
    print(f"\n=== {path} ===")
    doc = docx.Document(path)
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if any(k in t for k in ['พ211', 'พ221', 'พ231', 'พ311', 'พ321', 'พ331', 'รหัสวิชา พ', 'รหัส พ']):
            print(f"P{i}: {t}")
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            txt = " | ".join(c.text.strip().replace('\n', ' ') for c in row.cells)
            if any(k in txt for k in ['พ211', 'พ221', 'พ231', 'พ311', 'พ321', 'พ331']):
                print(f"T{t_idx} R{r_idx}: {txt[:120]}")

scan_pe_doc('school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/สุขศึกษา 69/5 ส่วนประกอบของหลักสูตร ฯ สุขศึกษา ม.ต้น มีน 11พ.ค.69.docx')
scan_pe_doc('school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/สุขศึกษา 69/5 ส่วนประกอบของหลักสูตร ฯ สุขศึกษาและพลศึกษา ม.ปลาย พ.ค. 2569.docx')
