# -*- coding: utf-8 -*-
import docx
import re
import os

files = [
    r"school_curriculum_raw\หลักสูตรกลุ่มสาระฯ 69\สังคม 69\ส่วนประกอบของหลักสูตร กลุ่มสาระสังคมศึกษา ม.ต้น 2569-2571.docx",
    r"school_curriculum_raw\หลักสูตรกลุ่มสาระฯ 69\สังคม 69\ส่วนประกอบของหลักสูตร กลุ่มสาระสังคมศึกษา ม.ปลาย 2569-2571.docx"
]

out = []

for file_path in files:
    out.append("="*80)
    out.append(f"FILE: {file_path}")
    out.append("="*80)
    doc = docx.Document(file_path)
    
    # Extract headings and paragraphs
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    out.append(f"Total non-empty paragraphs: {len(paragraphs)}")
    
    # Search for course codes (e.g. ส21101, ส31101, etc.)
    course_mentions = set()
    kpi_mentions = set()
    
    for p in paragraphs:
        # Check course codes
        matches = re.findall(r'[ส][0-9]{5}', p)
        for m in matches:
            course_mentions.add(m)
        # Check KPI codes (e.g. ส 1.1 ม.1/1, ส 5.1 ม.2/3)
        kpi_matches = re.findall(r'ส\s*[1-5]\.[1-2]\s*ม\.[1-6]/[0-9]+|ส\s*[1-5]\.[1-2]\s*ม\.[4-6]/[0-9]+', p)
        for km in kpi_matches:
            kpi_mentions.add(km)
            
    # Also search inside tables!
    table_kpi_mentions = set()
    for t_idx, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                matches = re.findall(r'[ส][0-9]{5}', text)
                for m in matches:
                    course_mentions.add(m)
                kpi_matches = re.findall(r'ส\s*[1-5]\.[1-2]\s*ม\.[1-6]/[0-9]+|ส\s*[1-5]\.[1-2]\s*ม\.[4-6]/[0-9]+', text)
                for km in kpi_matches:
                    table_kpi_mentions.add(km)

    out.append(f"Courses found: {sorted(course_mentions)}")
    out.append(f"KPIs in paragraphs: {len(kpi_mentions)}")
    out.append(f"KPIs in tables: {len(table_kpi_mentions)}")
    all_kpis = kpi_mentions.union(table_kpi_mentions)
    out.append(f"Total unique KPIs found: {len(all_kpis)}")
    out.append(f"KPI list: {sorted(all_kpis)}")

with open('school_social_curriculum_scan.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))

print("Scan completed.")
