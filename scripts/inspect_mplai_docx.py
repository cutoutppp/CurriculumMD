# -*- coding: utf-8 -*-
import docx
import re

doc = docx.Document(r'school_curriculum_raw\หลักสูตรกลุ่มสาระฯ 69\สังคม 69\ส่วนประกอบของหลักสูตร กลุ่มสาระสังคมศึกษา ม.ปลาย 2569-2571.docx')

out = []
out.append(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

# Search for course descriptions and indicator sections
courses_found = []
current_course = None

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    
    # Check course header
    # e.g., คำอธิบายรายวิชาพื้นฐาน ส31101 สังคมศึกษาฯ 1
    m_course = re.search(r'(คำอธิบายรายวิชา[^\n]*\b(ส\s*[0-9]{5})\b[^\n]*)', text)
    if not m_course:
        m_course = re.search(r'\b(ส\s*[0-9]{5})\b\s*รายวิชา[^\n]*', text)
    if not m_course:
        m_course = re.search(r'\b(ส\s*[0-9]{5})\b', text)
        
    if 'คำอธิบายรายวิชา' in text or (m_course and ('พื้นฐาน' in text or 'เพิ่มเติม' in text)):
        out.append(f"\n[P {i}] HEADER: {text}")
        
    # Search for indicators / results text
    if 'ตัวชี้วัด' in text or 'ผลการเรียนรู้' in text or 'รวมทั้งหมด' in text:
        out.append(f"  [P {i}] {text}")

# Also check tables
out.append("\n" + "="*50 + "\nTABLE SAMPLES\n" + "="*50)
for t_idx, table in enumerate(doc.tables):
    t_text = []
    has_kpi = False
    has_course = False
    for r in table.rows:
        row_str = " | ".join(c.text.strip().replace('\n', ' ') for c in r.cells)
        if any(k in row_str for k in ['ตัวชี้วัด', 'ผลการเรียนรู้', 'ส 1.', 'ส 2.', 'ส 3.', 'ส 4.', 'ส 5.', 'ส3']):
            has_kpi = True
        t_text.append(row_str)
    if has_kpi:
        out.append(f"\n--- TABLE {t_idx} ({len(table.rows)} rows) ---")
        for line in t_text[:10]:
            out.append("  " + line[:120])

with open('m_plai_curriculum_details.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))

print("M.Plai scan completed.")
