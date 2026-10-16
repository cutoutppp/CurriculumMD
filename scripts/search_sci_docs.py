import docx
import os
import re

def search_sci():
    paths = [
        'school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/วิทยาศาสตร์และ IS 69/รวมหลักสูตร ม.ต้น/คำอธิบายรายวิชา_ต้น_69.docx',
        'school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/วิทยาศาสตร์และ IS 69/รวมหลักสูตร ม.ต้น/โครงสร้าง_ต้น_69.docx',
        'school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/วิทยาศาสตร์และ IS 69/รวมหลักสูตร ม.ปลาย/คำอธิบายรายวิชาพื้นฐาน-รวม69.docx',
        'school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/วิทยาศาสตร์และ IS 69/รวมหลักสูตร ม.ปลาย/โครงสร้างรายวิชา-ุรวม69.docx',
    ]
    with open('sci_docs_summary.txt', 'w', encoding='utf-8') as out:
        for p in paths:
            if not os.path.exists(p):
                out.write(f"NOT FOUND: {p}\n")
                continue
            out.write(f"\n========================================\nFILE: {p}\n========================================\n")
            doc = docx.Document(p)
            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if any(k in text for k in ['ว 1.', 'ว 2.', 'ว 3.', 'ว 4.', 'ว2', 'ว3', 'คำอธิบายรายวิชา', 'รหัสวิชา', 'ตัวชี้วัด']):
                    out.write(f"P{i}: {text}\n")
            for t_idx, table in enumerate(doc.tables):
                for r_idx, row in enumerate(table.rows):
                    row_text = " | ".join(c.text.strip() for c in row.cells)
                    if any(k in row_text for k in ['ว 1.', 'ว 2.', 'ว 3.', 'ว 4.', 'ว2', 'ว3', 'รหัสวิชา', 'ตัวชี้วัด']):
                        out.write(f"T{t_idx} R{r_idx}: {row_text}\n")

search_sci()
print("sci_docs_summary.txt written")
