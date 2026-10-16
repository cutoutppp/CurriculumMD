import os
import docx
import re

target_files = []
for r, d, fs in os.walk('school_curriculum_raw'):
    for f in fs:
        if f.endswith('.docx') and 'สุข' in f:
            target_files.append(os.path.join(r, f))

with open('pe_deep_scan.txt', 'w', encoding='utf-8') as out:
    for fpath in target_files:
        out.write(f"\n========================================\nFILE: {fpath}\n========================================\n")
        try:
            doc = docx.Document(fpath)
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if any(k in t for k in ['พ 2', 'พ 3', 'พ2', 'พ3', 'คำอธิบายรายวิชา', 'ตัวชี้วัด', 'ผลการเรียนรู้', 'รหัส']):
                    out.write(f"P{i}: {t}\n")
            for t_idx, table in enumerate(doc.tables):
                for r_idx, row in enumerate(table.rows):
                    txt = " | ".join(c.text.strip().replace('\n', ' ') for c in row.cells)
                    if any(k in txt for k in ['พ211', 'พ221', 'พ231', 'พ311', 'พ321', 'พ331', 'ตัวชี้วัด', 'รหัส']):
                        out.write(f"T{t_idx} R{r_idx}: {txt}\n")
        except Exception as e:
            out.write(f"ERROR: {e}\n")

print("pe_deep_scan.txt completed")
