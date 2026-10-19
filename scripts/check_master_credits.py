# -*- coding: utf-8 -*-
import openpyxl
import sys
import re
import json

sys.stdout.reconfigure(encoding="utf-8")

excel_path = "school_curriculum_raw/หลักสูตรกลุ่มสาระฯ 69/ โครงสร้างหลักสูตร  โรงเรียนพัฒนานิคม(ปี2569-2571) แก้ไขรายวิชา.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

excel_courses = {}

for sname in ["2569-2571(ม.1,4)", "รายวิชาเพิ่มเติมเสรี ", "ก่อนแก้ 2569-2571(ม.1,4)"]:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column):
            val = str(ws.cell(r, c).value or "")
            m = re.search(r"([ก-ฮI]\s*\d{4,5})", val)
            if m:
                code = re.sub(r"\s+", "", m.group(1))
                nxt = ws.cell(r, c + 1).value
                if isinstance(nxt, (int, float)) and code not in excel_courses:
                    excel_courses[code] = float(nxt)

print(f"Total parsed from Excel: {len(excel_courses)}")

with open("school_courses_assessment_matrix.json", "r", encoding="utf-8") as f:
    matrix = json.load(f)

diffs = []
for code, cr in sorted(excel_courses.items()):
    if code in matrix:
        m_cr = float(matrix[code].get("credits", 0))
        if m_cr != cr:
            diffs.append((code, matrix[code].get("course_name"), m_cr, cr))

print(f"\nExcel vs Matrix Credit Diffs ({len(diffs)}):")
for code, name, m_cr, e_cr in diffs:
    print(f"  {code} ({name}): Matrix={m_cr} cr  vs  Excel={e_cr} cr")
