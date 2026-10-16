# -*- coding: utf-8 -*-
import docx
import re
import json

THAI_TO_ARABIC = str.maketrans('๐๑๒๓๔๕๖๗๘๙', '0123456789')
def clean(t):
    return t.translate(THAI_TO_ARABIC) if t else ''

def extract_all_basic_courses(m_ton_path, m_plai_path):
    results = {}
    
    for path, stage in [(m_ton_path, 'ม.ต้น'), (m_plai_path, 'ม.ปลาย')]:
        doc = docx.Document(path)
        for idx, p in enumerate(doc.paragraphs):
            t_orig = p.text.strip()
            t = clean(t_orig)
            
            # Match course heading in description: e.g. รหัส ส21101 or รหัสวิชา ส21101
            # Or รหัส ส... รายวิชา... or รหัส ส... วิชา...
            m = re.search(r'รหัส\s*(?:วิชา)?\s*(ส\s*[23]\d{4})\s*(?:รายวิชา|วิชา)?\s*([^\t\n\r]+)', t)
            if not m:
                continue
            code = m.group(1).replace(' ', '')
            name = m.group(2).strip()
            
            # Check context: is there 'คำอธิบายรายวิชาพื้นฐาน' or 'เวลาเรียน' nearby?
            is_desc = False
            for look in range(max(0, idx-4), min(len(doc.paragraphs), idx+5)):
                look_t = clean(doc.paragraphs[look].text)
                if any(k in look_t for k in ['คำอธิบายรายวิชาพื้นฐาน', 'คำอธิบายรายวิชา', 'เวลาเรียน', 'หน่วยกิต']):
                    is_desc = True
                    break
            if not is_desc:
                continue
                
            # If code is basic course code:
            if re.match(r'^ส[23][1-3]10[1-4]$', code):
                # Special check: in M.Ton, P1181 was mistyped as ส22102 in some places or ส22104:
                # Let's inspect the actual title text
                if 'ประวัติศาสตร์ 4' in name and code == 'ส22102':
                    code = 'ส22104'
                elif 'ประวัติศาสตร์ 3' in name and code == 'ส22104':
                    code = 'ส22102'
                    
                # Collect lines until next course or 'รวม ... ตัวชี้วัด'
                collected = [t_orig]
                for j in range(idx+1, min(len(doc.paragraphs), idx+45)):
                    nj_orig = doc.paragraphs[j].text.strip()
                    nj = clean(nj_orig)
                    if re.search(r'รหัส\s*(?:วิชา)?\s*ส\s*[23]\d{4}', nj) and any(k in nj for k in ['วิชา', 'กลุ่มสาระ']):
                        break
                    if nj_orig:
                        collected.append(nj_orig)
                    if re.search(r'รวม\s*\d+\s*(?:ตัวชี้วัด|ผลการเรียนรู้)', nj):
                        break
                        
                results[code] = {
                    'code': code,
                    'name': name,
                    'p_idx': idx,
                    'stage': stage,
                    'raw_lines': collected
                }
    return results

courses = extract_all_basic_courses(
    r'school_curriculum_raw\หลักสูตรกลุ่มสาระฯ 69\สังคม 69\ส่วนประกอบของหลักสูตร กลุ่มสาระสังคมศึกษา ม.ต้น 2569-2571.docx',
    r'school_curriculum_raw\หลักสูตรกลุ่มสาระฯ 69\สังคม 69\ส่วนประกอบของหลักสูตร กลุ่มสาระสังคมศึกษา ม.ปลาย 2569-2571.docx'
)

with open('extracted_22_courses.json', 'w', encoding='utf-8') as f:
    json.dump(courses, f, ensure_ascii=False, indent=2)

print(f'Total courses extracted: {len(courses)}')
for k in sorted(courses.keys()):
    print(k, courses[k]['name'], len(courses[k]['raw_lines']))
